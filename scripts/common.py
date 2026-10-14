# -*- coding: utf-8 -*-
"""
Outils partagés : configuration, chemins, lecture/écriture sûre des classeurs
(y compris chiffrés), détection des fichiers ouverts dans Excel, lecture des
produits, des achats et des ventes.

Compatible Excel 2016 : aucune fonction moderne n'est écrite dans les fichiers.
"""
from __future__ import annotations

import configparser
import datetime as dt
import io
import json
import logging
import os
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from logging.handlers import RotatingFileHandler
from pathlib import Path

import msoffcrypto
from openpyxl import Workbook, load_workbook

for _flux in (sys.stdout, sys.stderr):                 # console Windows : pas d'erreur sur ✔ ou é
    try:
        _flux.reconfigure(errors="replace")
    except Exception:
        pass

# --------------------------------------------------------------------------
# Chemins & configuration
# --------------------------------------------------------------------------
RACINE = Path(__file__).resolve().parent.parent          # ...\Projet_Stock
DOSSIER_DATA = RACINE / "data"                            # état interne
DOSSIER_LOGS = RACINE / "logs"
DOSSIER_SAUV = RACINE / "_sauvegardes"
DOSSIER_CONNAISSANCES = RACINE / "connaissances"          # base de connaissances du bot

MOIS_FR = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
           "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
JOURS_FR = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

# Noms des feuilles (référencés par les formules : ne pas renommer)
F_CATALOGUE = "Catalogue"
F_SAISIE_VENTES = "Saisie_ventes"
F_PARAMS = "Parametres"
F_SAISIE_ACHATS = "Saisie_achats"
F_PRODUITS_ACHATS = "Produits"

# Géométrie des feuilles de saisie (1re ligne de données)
LIGNE_DEBUT_SAISIE = 5
LIGNE_DEBUT_CATALOGUE = 5


@dataclass
class Config:
    annee: int
    nom_magasin: str
    f_source: Path
    f_vendeur: Path
    f_ventes: Path
    f_achats: Path
    f_admin: Path
    mdp_admin: str
    mdp_feuilles: str
    vendeurs: list[str]
    lignes_ventes: int
    lignes_achats: int
    database_url: str
    embed_provider: str
    embed_key: str
    embed_model: str
    embed_dim: int
    remise_max_pct: float = 7.0
    marge_min_pct: float = 15.0


def charger_config() -> Config:
    cp = configparser.ConfigParser(inline_comment_prefixes=(";",))
    cp.read(RACINE / "config.ini", encoding="utf-8")
    annee = cp.getint("general", "annee", fallback=dt.date.today().year)

    def f(cle, defaut):
        return RACINE / cp.get("fichiers", cle, fallback=defaut).format(annee=annee)

    # Les variables d'environnement priment (pratique pour ne pas écrire les secrets)
    return Config(
        annee=annee,
        nom_magasin=cp.get("general", "nom_magasin", fallback="Magasin"),
        f_source=f("source_produits", "Gestion_stock_corrige.xlsx"),
        f_vendeur=f("vendeur", "Stocks_vendeur.xlsx"),
        f_ventes=f("ventes", "Ventes_{annee}_Mois.xlsx"),
        f_achats=f("achats", "Achats_{annee}.xlsx"),
        f_admin=f("admin", "Stocks_admin.xlsx"),
        mdp_admin=os.environ.get("STOCK_MDP_ADMIN") or cp.get("securite", "mot_de_passe_admin"),
        mdp_feuilles=cp.get("securite", "mot_de_passe_feuilles", fallback="protect"),
        vendeurs=[v.strip() for v in cp.get("vendeurs", "noms").split(",") if v.strip()],
        lignes_ventes=cp.getint("saisie", "lignes_ventes", fallback=5000),
        lignes_achats=cp.getint("saisie", "lignes_achats", fallback=2000),
        database_url=os.environ.get("NEON_DATABASE_URL") or cp.get("neon", "database_url", fallback="").strip(),
        embed_provider=(os.environ.get("EMBED_PROVIDER") or cp.get("embeddings", "provider", fallback="none")).strip().lower(),
        embed_key=os.environ.get("EMBED_API_KEY") or cp.get("embeddings", "api_key", fallback="").strip(),
        embed_model=os.environ.get("EMBED_MODEL") or cp.get("embeddings", "model", fallback="").strip(),
        embed_dim=int(os.environ.get("EMBED_DIM") or cp.get("embeddings", "dimension", fallback="1024")),
        remise_max_pct=cp.getfloat("negociation", "remise_max_pourcent", fallback=7.0),
        marge_min_pct=cp.getfloat("negociation", "marge_min_pourcent", fallback=15.0),
    )


def prix_plancher(p: "Produit", cfg: "Config") -> float:
    """Prix le plus bas que le bot peut accorder en négociation (jamais montré au client).
    Colonne « Prix minimum » si remplie ; sinon le plus élevé de :
      prix de vente - remise max (%)   et   prix d'achat + marge minimale (%).
    Arrondi au millier supérieur, jamais au-dessus du prix de vente."""
    import math
    if p.prix_min and p.prix_min > 0:
        plancher = p.prix_min
    else:
        plancher = max(p.prix_vente * (1 - cfg.remise_max_pct / 100),
                       p.prix_achat * (1 + cfg.marge_min_pct / 100) if p.prix_achat else 0)
        plancher = math.ceil(plancher / 1000) * 1000
    return min(plancher, p.prix_vente)


def init_logs(nom: str = "sync") -> logging.Logger:
    DOSSIER_LOGS.mkdir(exist_ok=True)
    log = logging.getLogger(nom)
    if not log.handlers:
        log.setLevel(logging.INFO)
        fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S")
        fh = RotatingFileHandler(DOSSIER_LOGS / f"{nom}.log", maxBytes=1_000_000, backupCount=5, encoding="utf-8")
        fh.setFormatter(fmt)
        log.addHandler(fh)
        if sys.stderr is not None:                      # absent avec pythonw (tâche planifiée)
            sh = logging.StreamHandler()
            sh.setFormatter(fmt)
            log.addHandler(sh)
    return log


# --------------------------------------------------------------------------
# Fichiers ouverts / écriture atomique / chiffrement
# --------------------------------------------------------------------------
def est_ouvert_dans_excel(chemin: Path) -> bool:
    """Vrai si le classeur est ouvert dans Excel (fichier ~$ présent ou verrou Windows)."""
    if not chemin.exists():
        return False
    proprietaire = chemin.with_name("~$" + chemin.name)
    if proprietaire.exists():
        # Excel garde ce petit fichier ouvert (verrouillé) tant que le classeur est ouvert.
        # S'il n'est pas verrouillé, c'est un reste d'Excel fermé brutalement : on le supprime.
        try:
            with open(proprietaire, "r+b"):
                pass
            os.remove(proprietaire)
            logging.getLogger("sync").info("Fichier verrou orphelin supprimé : %s", proprietaire.name)
        except PermissionError:
            return True                       # verrouillé : le classeur est vraiment ouvert
        except OSError:
            pass
    try:
        with open(chemin, "r+b"):
            pass
        return False
    except PermissionError:
        return True
    except OSError:
        return False


def est_chiffre(chemin: Path) -> bool:
    try:
        with open(chemin, "rb") as fh:
            return msoffcrypto.OfficeFile(fh).is_encrypted()
    except Exception:
        return False


def ouvrir_classeur(chemin: Path, mot_de_passe: str | None = None, **kwargs):
    """Ouvre un classeur, chiffré ou non. Lecture possible même s'il est ouvert dans Excel."""
    with open(chemin, "rb") as fh:
        brut = fh.read()
    flux = io.BytesIO(brut)
    try:
        office = msoffcrypto.OfficeFile(io.BytesIO(brut))
        chiffre = office.is_encrypted()
    except Exception:
        chiffre = False                      # xlsx classique (non OLE)
    if chiffre:
        if not mot_de_passe:
            raise RuntimeError(f"{chemin.name} est chiffré et aucun mot de passe n'est fourni")
        try:
            office.load_key(password=mot_de_passe)
            flux = io.BytesIO()
            office.decrypt(flux)
            flux.seek(0)
        except Exception as e:
            raise RuntimeError(f"Mot de passe incorrect pour {chemin.name} (vérifiez config.ini)") from e
    return load_workbook(flux, **kwargs)


class FichierVerrouille(Exception):
    pass


def enregistrer_classeur(wb: Workbook, chemin: Path, mot_de_passe: str | None = None) -> None:
    """Enregistre via un fichier temporaire puis remplacement atomique.
    Lève FichierVerrouille si le fichier est ouvert dans Excel (jamais d'écrasement partiel)."""
    if est_ouvert_dans_excel(chemin):
        raise FichierVerrouille(chemin.name)
    tampon = io.BytesIO()
    wb.save(tampon)
    donnees = tampon.getvalue()
    if mot_de_passe:
        from msoffcrypto.format.ooxml import OOXMLFile
        chiffre = io.BytesIO()
        OOXMLFile(io.BytesIO(donnees)).encrypt(mot_de_passe, chiffre)
        donnees = chiffre.getvalue()
    fd, tmp = tempfile.mkstemp(prefix=".~sync_", suffix=".xlsx", dir=chemin.parent)
    try:
        with os.fdopen(fd, "wb") as out:
            out.write(donnees)
        try:
            os.chmod(tmp, 0o644)
        except OSError:
            pass
        try:
            os.replace(tmp, chemin)
        except PermissionError as e:
            raise FichierVerrouille(chemin.name) from e
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def lire_json(chemin: Path, defaut):
    try:
        return json.loads(chemin.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return defaut


def ecrire_json(chemin: Path, donnees) -> None:
    chemin.parent.mkdir(exist_ok=True)
    tmp = chemin.with_suffix(".tmp")
    tmp.write_text(json.dumps(donnees, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    os.replace(tmp, chemin)


# --------------------------------------------------------------------------
# Normalisation
# --------------------------------------------------------------------------
def normaliser(texte) -> str:
    s = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def nombre(v, defaut=0.0) -> float:
    if v is None or v == "":
        return defaut
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace("\xa0", "").replace(" ", "")
    s = re.sub(r"(?i)(fcfa|xaf|cfa|f)$", "", s).replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return defaut


def en_date(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:            # numéro de série Excel
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def fcfa(montant: float) -> str:
    return f"{montant:,.0f}".replace(",", " ") + " FCFA"


# --------------------------------------------------------------------------
# Produits (source = Gestion_stock_corrige.xlsx)
# --------------------------------------------------------------------------
@dataclass
class Produit:
    code: str
    designation: str
    categorie: str = ""
    marque: str = ""
    prix_achat: float = 0.0
    prix_vente: float = 0.0
    stock_initial: float = 0.0
    seuil: float = 2.0
    description: str = ""
    photo_url: str = ""
    usages: str = ""
    prix_min: float = 0.0          # prix plancher de négociation (0 = calculé automatiquement)


SYNONYMES = {
    "code": ["code", "reference", "ref", "code produit", "id", "sku", "code article"],
    "designation": ["designation", "produit", "nom", "article", "libelle", "nom produit", "designation produit"],
    "categorie": ["categorie", "famille", "type", "rayon"],
    "marque": ["marque", "fabricant"],
    "prix_achat": ["prix achat", "pa", "cout", "prix d achat", "cout unitaire", "prix achat unitaire", "pu achat"],
    "prix_vente": ["prix vente", "pv", "prix", "prix de vente", "prix unitaire", "pu vente", "prix vente unitaire"],
    "stock_initial": ["stock initial", "stock", "quantite", "qte", "stock depart", "stock de depart", "quantite initiale", "stock actuel"],
    "seuil": ["seuil", "seuil alerte", "stock minimum", "stock min", "alerte", "seuil d alerte"],
    "description": ["description", "caracteristiques", "specifications", "details"],
    "photo_url": ["photo", "image", "photo url", "url photo", "lien photo", "photos"],
    "usages": ["usages", "usage", "profil", "ideal pour", "conseille pour"],
    "prix_min": ["prix minimum", "prix plancher", "prix min", "prix negociable", "prix minimum negociation",
                 "prix mini"],
}


def _associer_entetes(ligne) -> dict[str, int]:
    trouve: dict[str, int] = {}
    entetes = [normaliser(c) for c in ligne]
    entetes = [re.sub(r"\b(fcfa|xaf|cfa|en|u)\b", "", e).strip() for e in entetes]
    for champ, syns in SYNONYMES.items():
        for i, e in enumerate(entetes):
            if e in syns and i not in trouve.values():
                trouve[champ] = i
                break
    for champ, syns in SYNONYMES.items():           # 2e passe : correspondance partielle
        if champ in trouve:
            continue
        for i, e in enumerate(entetes):
            if e and i not in trouve.values() and any(e.startswith(s) for s in syns if len(s) > 3):
                trouve[champ] = i
                break
    return trouve


def charger_produits(cfg: Config) -> list[Produit]:
    """Lit la liste des produits depuis le fichier source, quelle que soit la feuille
    et l'ordre des colonnes (détection par les en-têtes)."""
    wb = ouvrir_classeur(cfg.f_source, cfg.mdp_admin, data_only=True, read_only=True)
    meilleur: tuple[int, list[Produit]] = (0, [])
    for ws in wb.worksheets:
        lignes = list(ws.iter_rows(values_only=True))
        for idx, ligne in enumerate(lignes[:15]):
            cols = _associer_entetes(ligne)
            if "designation" not in cols or not ({"prix_vente", "code"} & cols.keys()):
                continue
            produits = []
            for r in lignes[idx + 1:]:
                val = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
                nom = val("designation")
                if not nom or str(nom).strip().lower().startswith("total"):
                    continue
                code = str(val("code") or "").strip() or f"P{len(produits) + 1:03d}"
                produits.append(Produit(
                    code=code.upper(),
                    designation=str(nom).strip(),
                    categorie=str(val("categorie") or "Divers").strip(),
                    marque=str(val("marque") or "").strip(),
                    prix_achat=nombre(val("prix_achat")),
                    prix_vente=nombre(val("prix_vente")),
                    stock_initial=nombre(val("stock_initial")),
                    seuil=nombre(val("seuil"), 2.0),
                    description=str(val("description") or "").strip(),
                    photo_url=str(val("photo_url") or "").strip(),
                    usages=str(val("usages") or "").strip(),
                    prix_min=nombre(val("prix_min")),
                ))
            if len(produits) > meilleur[0]:
                meilleur = (len(produits), produits)
            break
    wb.close()
    if not meilleur[1]:
        raise RuntimeError(
            f"Aucun produit trouvé dans {cfg.f_source.name}. Il faut au minimum les colonnes "
            "'Désignation' et 'Prix de vente' (ou 'Code').")
    vus, uniques = set(), []
    for p in meilleur[1]:
        if p.code in vus:
            p.code = f"{p.code}-{len(uniques) + 1}"
        vus.add(p.code)
        uniques.append(p)
    return uniques


# --------------------------------------------------------------------------
# Achats & ventes
# --------------------------------------------------------------------------
@dataclass
class Achat:
    ligne: int
    date: dt.date
    code: str
    quantite: float
    prix_unitaire: float
    fournisseur: str = ""
    facture: str = ""

    @property
    def total(self):
        return self.quantite * self.prix_unitaire


@dataclass
class Vente:
    ligne: int
    date: dt.date
    code: str
    quantite: float
    prix_unitaire: float
    vendeur: str
    client: str = ""
    date_estimee: bool = False
    designation: str = ""
    cout_unitaire: float = 0.0

    @property
    def total(self):
        return self.quantite * self.prix_unitaire

    @property
    def marge(self):
        return self.quantite * (self.prix_unitaire - self.cout_unitaire)


@dataclass
class Anomalie:
    fichier: str
    ligne: int
    message: str


def lire_achats(cfg: Config, produits: dict[str, Produit], anomalies: list[Anomalie]) -> list[Achat]:
    if not cfg.f_achats.exists():
        return []
    wb = ouvrir_classeur(cfg.f_achats, cfg.mdp_admin, data_only=True, read_only=True)
    ws = wb[F_SAISIE_ACHATS]
    achats = []
    for n, r in enumerate(ws.iter_rows(min_row=LIGNE_DEBUT_SAISIE, values_only=True), start=LIGNE_DEBUT_SAISIE):
        # Colonnes : A Date | B Code | C Désignation | D Qté | E PA unitaire | F Fournisseur | G N° facture | H Total
        date, code, _, qte, pu, fourn, fact = (list(r) + [None] * 8)[:7]
        code = str(code or "").strip().upper()
        if not code and not qte:
            continue
        q = nombre(qte)
        if code not in produits:
            anomalies.append(Anomalie(cfg.f_achats.name, n, f"Code produit inconnu « {code} »"))
            continue
        if q <= 0:
            anomalies.append(Anomalie(cfg.f_achats.name, n, "Quantité manquante ou nulle"))
            continue
        d = en_date(date)
        if d is None:
            anomalies.append(Anomalie(cfg.f_achats.name, n, "Date manquante ou illisible (achat ignoré)"))
            continue
        achats.append(Achat(n, d, code, q, nombre(pu) or produits[code].prix_achat,
                            str(fourn or "").strip(), str(fact or "").strip()))
    wb.close()
    return achats


def lire_ventes(cfg: Config, produits: dict[str, Produit], anomalies: list[Anomalie]) -> list[Vente]:
    """Lit la saisie des vendeurs. Le prix catalogue est « figé » au premier passage
    du script (journal), pour qu'un changement de prix ne réécrive pas l'historique."""
    journal_path = DOSSIER_DATA / "journal_ventes.json"
    journal: dict = lire_json(journal_path, {})
    nouveau: dict = {}
    aujourd_hui = dt.date.today()

    wb = ouvrir_classeur(cfg.f_vendeur, data_only=True, read_only=True)
    ws = wb[F_SAISIE_VENTES]
    ventes = []
    for n, r in enumerate(ws.iter_rows(min_row=LIGNE_DEBUT_SAISIE, values_only=True), start=LIGNE_DEBUT_SAISIE):
        # A N° | B Date | C Code | D Désignation | E Qté | F Prix catalogue | G Prix pratiqué | H Total | I Vendeur | J Client
        _, date, code, _, qte, _, prix_prat, _, vendeur, client = (list(r) + [None] * 11)[:10]
        code = str(code or "").strip().upper()
        if not code and not qte:
            continue
        cle = str(n)
        q = nombre(qte)
        if code not in produits:
            anomalies.append(Anomalie(cfg.f_vendeur.name, n, f"Code produit inconnu « {code} »"))
            continue
        if q <= 0:
            anomalies.append(Anomalie(cfg.f_vendeur.name, n, "Quantité manquante ou nulle"))
            continue

        ancien = journal.get(cle, {})
        if ancien.get("code") == code:
            prix_fige = ancien["prix"]
            vu_le = ancien["vu_le"]
        else:
            prix_fige = produits[code].prix_vente
            vu_le = aujourd_hui.isoformat()
        nouveau[cle] = {"code": code, "prix": prix_fige, "vu_le": vu_le}

        d = en_date(date)
        estimee = False
        if d is None:
            d = dt.date.fromisoformat(vu_le)
            estimee = True
            anomalies.append(Anomalie(cfg.f_vendeur.name, n,
                                      f"Date manquante : date de saisie détectée utilisée ({d:%d/%m/%Y})"))
        pp = nombre(prix_prat)
        vend = str(vendeur or "").strip() or "Non renseigné"
        if vend == "Non renseigné":
            anomalies.append(Anomalie(cfg.f_vendeur.name, n, "Vendeur non renseigné"))
        p = produits[code]
        ventes.append(Vente(n, d, code, q, pp if pp > 0 else prix_fige, vend,
                            str(client or "").strip(), estimee, p.designation, p.prix_achat))
    wb.close()
    ecrire_json(journal_path, nouveau)
    return ventes


@dataclass
class EtatStock:
    produit: Produit
    entrees: float = 0.0
    sorties: float = 0.0
    ca: float = 0.0
    marge: float = 0.0
    dernier_achat: Achat | None = None
    derniere_vente: Vente | None = None

    @property
    def stock(self):
        return self.produit.stock_initial + self.entrees - self.sorties

    @property
    def statut(self):
        if self.stock <= 0:
            return "Rupture"
        if self.stock <= self.produit.seuil:
            return "Stock faible"
        return "Disponible"


def calculer_stocks(produits: list[Produit], achats: list[Achat], ventes: list[Vente]) -> dict[str, EtatStock]:
    etats = {p.code: EtatStock(p) for p in produits}
    for a in sorted(achats, key=lambda x: x.date):
        e = etats[a.code]
        e.entrees += a.quantite
        e.dernier_achat = a
    for v in sorted(ventes, key=lambda x: x.date):
        e = etats[v.code]
        e.sorties += v.quantite
        e.ca += v.total
        e.marge += v.marge
        e.derniere_vente = v
    return etats


# --------------------------------------------------------------------------
# Verrou d'exécution (évite 2 synchros simultanées)
# --------------------------------------------------------------------------
class VerrouExecution:
    def __init__(self, nom="sync"):
        self.chemin = DOSSIER_DATA / f"{nom}.lock"

    def __enter__(self):
        DOSSIER_DATA.mkdir(exist_ok=True)
        if self.chemin.exists():
            age = dt.datetime.now().timestamp() - self.chemin.stat().st_mtime
            if age < 600:
                raise RuntimeError("Une synchronisation est déjà en cours")
        self.chemin.write_text(str(os.getpid()))
        return self

    def __exit__(self, *exc):
        try:
            self.chemin.unlink()
        except FileNotFoundError:
            pass
