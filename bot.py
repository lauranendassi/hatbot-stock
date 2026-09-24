# -*- coding: utf-8 -*-
"""Bot Messenger WeloobeAI — Version IA Groq robuste anti-hallucination.
   L'IA ne propose JAMAIS de produits absents du catalogue.
"""
import os
import json
import datetime
import unicodedata
import traceback
import requests
import psycopg2
import psycopg2.extras
from flask import Flask, request, jsonify
from openai import OpenAI

# ====================================================================
# CONFIG
# ====================================================================
PAGE_ACCESS_TOKEN = os.environ.get("PAGE_ACCESS_TOKEN", "")
VERIFY_TOKEN = os.environ.get("VERIFY_TOKEN", "weloobe_verify_2026_secure")
DATABASE_URL = os.environ.get("DATABASE_URL", "")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")

# Modeles Groq par ordre de preference (fallback automatique)
MODELES_GROQ = [
    "openai/gpt-oss-120b",
    "openai/gpt-oss-20b",
    "llama-3.1-8b-instant",
    "llama-3.1-70b-versatile",
]

app = Flask(__name__)

client_ia = None
if GROQ_API_KEY:
    try:
        client_ia = OpenAI(
            base_url="https://api.groq.com/openai/v1",
            api_key=GROQ_API_KEY
        )
    except Exception as e:
        print("Erreur init Groq :", e)


# ====================================================================
# UTILITAIRES
# ====================================================================
def normaliser(s):
    try:
        s = (s or "").lower()
        s = unicodedata.normalize("NFD", s)
        s = "".join(c for c in s if unicodedata.category(c) != "Mn")
        for ch in "—–-_,.;:!?()[]'\"/\\":
            s = s.replace(ch, " ")
        return " ".join(s.split())
    except Exception:
        return ""


def f(n):
    try:
        return "{:,}".format(int(n)).replace(",", " ") + " FCFA"
    except Exception:
        return "0 FCFA"


def db():
    if not DATABASE_URL:
        raise Exception("DATABASE_URL non configuree")
    return psycopg2.connect(
        DATABASE_URL,
        cursor_factory=psycopg2.extras.RealDictCursor,
        connect_timeout=10
    )


def log(prefixe, message):
    try:
        print("[{}] {}".format(prefixe, str(message)[:500]))
    except Exception:
        pass


# ====================================================================
# OUTILS POUR L'IA
# ====================================================================
def chercher_produits(requete="", budget_max=0, categorie=""):
    """Cherche des produits. Ne plante jamais.
       Retourne toujours une liste avec message explicite si vide."""
    try:
        conn = db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM produits")
        tous = cur.fetchall()
        cur.close()
        conn.close()
    except Exception as e:
        log("ERR", "chercher_produits DB: " + str(e))
        return [{"erreur": "Base momentanement indisponible"}]

    try:
        resultats = []
        tokens = normaliser(requete).split()
        try:
            budget_max = float(budget_max) if budget_max else 0
        except Exception:
            budget_max = 0

        for p in tous:
            try:
                hay = normaliser((p.get("nom") or "") + " " + (p.get("categorie") or ""))
                score = sum(len(t) for t in tokens if t in hay)
                if categorie and categorie.lower() not in (p.get("categorie") or "").lower():
                    continue
                if budget_max and (p.get("prix_vente") or 0) > budget_max:
                    continue
                if not tokens or score > 0:
                    resultats.append({
                        "sku": p.get("sku"),
                        "nom": p.get("nom"),
                        "categorie": p.get("categorie"),
                        "prix": p.get("prix_vente"),
                        "score": score
                    })
            except Exception:
                continue

        resultats.sort(key=lambda x: -x.get("score", 0))
        top = resultats[:5]

        # Cas vide : renvoyer un message CLAIR et NON EQUIVOQUE a l'IA
        if not top:
            prix_min = min((p.get("prix_vente") or 0) for p in tous) if tous else 0
            return [{
                "aucun_resultat": True,
                "message": "AUCUN PRODUIT du catalogue ne correspond a cette recherche.",
                "budget_demande": budget_max,
                "categorie_demandee": categorie,
                "prix_minimum_catalogue": prix_min,
                "instruction": "Dis honnetement au client qu'aucun produit ne correspond. Propose UNIQUEMENT d'elargir le budget ou de voir une autre categorie REELLE du catalogue (Ecran, Composant, Accessoire). N'invente AUCUN produit."
            }]

        return top
    except Exception as e:
        log("ERR", "chercher_produits traitement: " + str(e))
        return [{"erreur": "Erreur lors de la recherche"}]


def verifier_stock(sku):
    """Verifie le stock d'un produit. Ne plante jamais."""
    try:
        conn = db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM produits WHERE sku = %s", (sku,))
        p = cur.fetchone()
        cur.execute("SELECT COALESCE(SUM(quantite), 0) AS q FROM ventes WHERE sku = %s", (sku,))
        vendu = cur.fetchone()
        cur.close()
        conn.close()
        if not p:
            return {"erreur": "Produit inconnu dans le catalogue"}
        stock = (p.get("stock_initial") or 0) - (vendu.get("q") or 0)
        return {
            "sku": sku,
            "nom": p.get("nom"),
            "stock": stock,
            "prix": p.get("prix_vente"),
            "disponible": stock > 0
        }
    except Exception as e:
        log("ERR", "verifier_stock: " + str(e))
        return {"erreur": "Impossible de verifier le stock"}


def lister_catalogue():
    """Liste tout le catalogue. Ne plante jamais."""
    try:
        conn = db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM produits ORDER BY categorie, nom")
        produits = cur.fetchall()
        cur.close()
        conn.close()
        if not produits:
            return [{"info": "Catalogue vide"}]
        return [{
            "sku": p.get("sku"),
            "nom": p.get("nom"),
            "categorie": p.get("categorie"),
            "prix": p.get("prix_vente")
        } for p in produits]
    except Exception as e:
        log("ERR", "lister_catalogue: " + str(e))
        return [{"erreur": "Catalogue indisponible"}]


OUTILS = [
    {
        "type": "function",
        "function": {
            "name": "chercher_produits",
            "description": "Cherche des produits dans le catalogue WeloobeAI. Retourne UNIQUEMENT les produits existants. Si aucun ne correspond, indique-le clairement.",
            "parameters": {
                "type": "object",
                "properties": {
                    "requete": {"type": "string", "description": "Mots-cles (ex: 'pc portable', 'ecran')"},
                    "budget_max": {"type": "number", "description": "Budget maximum en FCFA (0 si non specifie)"},
                    "categorie": {"type": "string", "description": "Categorie EXACTE parmi : PC portable, PC fixe, Ecran, Composant, Accessoire"}
                }
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "verifier_stock",
            "description": "Verifie la disponibilite d'un produit par son SKU",
            "parameters": {
                "type": "object",
                "properties": {"sku": {"type": "string"}},
                "required": ["sku"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "lister_catalogue",
            "description": "Liste TOUS les produits du catalogue WeloobeAI",
            "parameters": {"type": "object", "properties": {}}
        }
    }
]


# ====================================================================
# HISTORIQUE
# ====================================================================
def charger_historique(psid, limite=10):
    try:
        conn = db()
        cur = conn.cursor()
        cur.execute(
            "SELECT role, contenu FROM messages WHERE psid = %s "
            "ORDER BY timestamp DESC LIMIT %s",
            (psid, limite)
        )
        rows = cur.fetchall()
        cur.close()
        conn.close()
        rows = list(reversed(rows))
        return [{"role": r.get("role") or "user", "content": r.get("contenu") or ""} for r in rows]
    except Exception as e:
        log("WARN", "charger_historique: " + str(e))
        return []


def enregistrer_message(psid, role, contenu):
    try:
        conn = db()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO messages (psid, role, contenu, timestamp) VALUES (%s, %s, %s, %s)",
            (psid, role, contenu or "", datetime.datetime.now().isoformat())
        )
        conn.commit()
        cur.close()
        conn.close()
    except Exception as e:
        log("WARN", "enregistrer_message: " + str(e))


# ====================================================================
# SYSTEM PROMPT — ANTI-HALLUCINATION
# ====================================================================
SYSTEM_PROMPT = """Tu es l'assistant commercial de WeloobeAI, magasin de materiel informatique a Yaounde (Cameroun).

=== REGLE ABSOLUE ET NON NEGOCIABLE ===
Tu ne dois JAMAIS inventer, suggerer, ou mentionner un produit, un modele, une marque, une categorie ou un prix qui ne vient PAS d'un appel d'outil.

INTERDICTIONS STRICTES :
- Ne JAMAIS mentionner : Chromebook, tablette, iPad, reconditionne, occasion, pack etudiant, partenaire externe, produits d'autres marques non listes.
- Ne JAMAIS inventer de prix.
- Ne JAMAIS proposer un produit si l'outil chercher_produits a retourne "aucun_resultat".

SI L'OUTIL RETOURNE "aucun_resultat" :
Tu dois dire HONNETEMENT au client :
1. Qu'aucun produit du catalogue ne correspond a sa recherche
2. Le prix du produit le moins cher du catalogue (fourni dans le champ "prix_minimum_catalogue")
3. Proposer UNIQUEMENT ces options REELLES :
   - Elargir le budget
   - Voir une autre categorie REELLE du catalogue : Ecran, Composant, Accessoire
   - Laisser ses coordonnees pour etre rappele

=== TON ROLE ===
- Accueillir chaleureusement les clients
- Comprendre leur besoin
- Proposer UNIQUEMENT les produits retournes par chercher_produits
- Donner les prix EXACTEMENT tels qu'ils viennent de la base
- Verifier la disponibilite avec verifier_stock

=== STYLE ===
- Naturel, amical, professionnel
- Emojis avec moderation
- Reponds toujours en francais

=== CATEGORIES REELLES DU CATALOGUE ===
- PC portable
- PC fixe
- Ecran
- Composant
- Accessoire"""


def appeler_ia(messages, avec_outils=True):
    """Appelle Groq avec fallback sur plusieurs modeles."""
    if not client_ia:
        return None

    for modele in MODELES_GROQ:
        try:
            kwargs = {
                "model": modele,
                "messages": messages,
                "temperature": 0.5,
                "max_tokens": 500
            }
            if avec_outils:
                kwargs["tools"] = OUTILS
                kwargs["tool_choice"] = "auto"

            response = client_ia.chat.completions.create(**kwargs)
            log("IA", "Modele utilise : " + modele)
            return response
        except Exception as e:
            log("WARN", "Modele {} echoue : {}".format(modele, str(e)[:200]))
            continue
    return None


def executer_outil(nom, args):
    """Execute un outil avec gestion d'erreur."""
    try:
        if nom == "chercher_produits":
            return chercher_produits(
                args.get("requete", ""),
                args.get("budget_max", 0),
                args.get("categorie", "")
            )
        elif nom == "verifier_stock":
            return verifier_stock(args.get("sku", ""))
        elif nom == "lister_catalogue":
            return lister_catalogue()
        return {"erreur": "outil inconnu"}
    except Exception as e:
        log("ERR", "executer_outil: " + str(e))
        return {"erreur": "Erreur execution outil"}


def repondre_avec_ia(psid, message_client):
    """Genere une reponse. Ne plante jamais."""

    if not client_ia:
        return "Bonjour ! Je suis l'assistant WeloobeAI. Le service est en cours de configuration, merci de reessayer dans quelques instants."

    try:
        historique = charger_historique(psid, limite=10)
        messages = [{"role": "system", "content": SYSTEM_PROMPT}]
        messages.extend(historique)
        messages.append({"role": "user", "content": message_client})

        # Premier appel
        response = appeler_ia(messages, avec_outils=True)
        if not response:
            return "Je rencontre un souci technique. Pouvez-vous reformuler ?"

        try:
            msg = response.choices[0].message
        except Exception:
            return "Je n'ai pas bien compris. Pouvez-vous reformuler ?"

        # Appel d'outils
        tool_calls = getattr(msg, "tool_calls", None)
        if tool_calls:
            messages.append({
                "role": "assistant",
                "content": msg.content or "",
                "tool_calls": [
                    {
                        "id": tc.id,
                        "type": "function",
                        "function": {
                            "name": tc.function.name,
                            "arguments": tc.function.arguments
                        }
                    } for tc in tool_calls
                ]
            })

            for tool_call in tool_calls:
                try:
                    nom_outil = tool_call.function.name
                    args_str = tool_call.function.arguments or "{}"
                    try:
                        args = json.loads(args_str)
                    except Exception:
                        args = {}

                    log("OUTIL", "{} ({})".format(nom_outil, args))
                    resultat = executer_outil(nom_outil, args)

                    messages.append({
                        "role": "tool",
                        "tool_call_id": tool_call.id,
                        "content": json.dumps(resultat, ensure_ascii=False, default=str)
                    })
                except Exception as e:
                    log("ERR", "traitement tool_call: " + str(e))
                    continue

            # Deuxieme appel pour formuler la reponse finale
            response2 = appeler_ia(messages, avec_outils=False)
            if response2:
                try:
                    contenu = response2.choices[0].message.content
                    if contenu:
                        return contenu
                except Exception:
                    pass

            return "J'ai trouve des produits mais je n'arrive pas a formuler la reponse. Reformulez svp."

        # Reponse directe
        return msg.content or "Je n'ai pas bien compris."

    except Exception as e:
        log("ERR", "repondre_avec_ia: " + str(e))
        log("ERR", traceback.format_exc()[:500])
        return "Je rencontre un souci technique. Pouvez-vous reformuler votre demande ?"


def envoyer_message(psid, texte):
    """Envoie un message a Messenger. Ne plante jamais."""
    if not PAGE_ACCESS_TOKEN:
        log("ERR", "PAGE_ACCESS_TOKEN manquant")
        return False
    try:
        url = "https://graph.facebook.com/v20.0/me/messages"
        params = {"access_token": PAGE_ACCESS_TOKEN}
        texte = (texte or "")[:1900]
        payload = {"recipient": {"id": psid}, "message": {"text": texte}}
        r = requests.post(url, params=params, json=payload, timeout=10)
        log("ENVOI", "{} - {}".format(r.status_code, texte[:60]))
        return r.status_code == 200
    except Exception as e:
        log("ERR", "envoyer_message: " + str(e))
        return False


# ====================================================================
# WEBHOOK
# ====================================================================
@app.route("/", methods=["GET"])
def accueil():
    return jsonify({
        "status": "ok",
        "bot": "WeloobeAI Chatbot IA",
        "ia": "connectee" if client_ia else "non configuree",
        "db": "connectee" if DATABASE_URL else "non configuree"
    })


@app.route("/webhook", methods=["GET"])
def webhook_verification():
    try:
        mode = request.args.get("hub.mode")
        token = request.args.get("hub.verify_token")
        challenge = request.args.get("hub.challenge")
        if mode == "subscribe" and token == VERIFY_TOKEN:
            log("WEBHOOK", "Verification OK")
            return challenge or "", 200
        log("WEBHOOK", "Verification echouee")
        return "Forbidden", 403
    except Exception as e:
        log("ERR", "webhook_verification: " + str(e))
        return "Forbidden", 403


@app.route("/webhook", methods=["POST"])
def webhook_reception():
    """Recoit les evenements. Ne plante JAMAIS."""
    try:
        data = request.get_json(silent=True) or {}
        if data.get("object") != "page":
            return "Not a page event", 404

        for entry in data.get("entry", []):
            for event in entry.get("messaging", []):
                try:
                    psid = event.get("sender", {}).get("id")
                    texte = event.get("message", {}).get("text", "")
                    if not psid or not texte:
                        continue

                    log("RECU", "{} : {}".format(psid, texte))
                    enregistrer_message(psid, "user", texte)

                    reponse = repondre_avec_ia(psid, texte)
                    if not reponse:
                        reponse = "Je n'ai pas de reponse pour le moment."

                    enregistrer_message(psid, "assistant", reponse)
                    envoyer_message(psid, reponse)

                except Exception as e:
                    log("ERR", "traitement event: " + str(e))
                    log("ERR", traceback.format_exc()[:500])
                    continue

        return "EVENT_RECEIVED", 200

    except Exception as e:
        log("ERR", "webhook_reception global: " + str(e))
        return "EVENT_RECEIVED", 200


# ====================================================================
# INIT BASE
# ====================================================================
def init_database():
    """Cree les tables. Ne plante jamais."""
    if not DATABASE_URL:
        log("WARN", "Pas de DATABASE_URL, base non initialisee")
        return

    try:
        conn = db()
        cur = conn.cursor()

        cur.execute("""
            CREATE TABLE IF NOT EXISTS produits (
                sku TEXT PRIMARY KEY,
                nom TEXT NOT NULL,
                categorie TEXT,
                prix_achat INTEGER,
                prix_vente INTEGER,
                stock_initial INTEGER,
                seuil INTEGER
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS clients (
                psid TEXT PRIMARY KEY,
                derniere_interaction TEXT
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS messages (
                id SERIAL PRIMARY KEY,
                psid TEXT,
                role TEXT,
                contenu TEXT,
                timestamp TEXT
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS ventes (
                id SERIAL PRIMARY KEY,
                date TEXT,
                sku TEXT,
                quantite INTEGER,
                client TEXT
            )
        """)

        for alter in [
            "ALTER TABLE messages ADD COLUMN IF NOT EXISTS role TEXT",
            "ALTER TABLE messages ADD COLUMN IF NOT EXISTS timestamp TEXT",
            "ALTER TABLE messages ADD COLUMN IF NOT EXISTS psid TEXT",
            "ALTER TABLE messages ADD COLUMN IF NOT EXISTS contenu TEXT",
        ]:
            try:
                cur.execute(alter)
            except Exception as e:
                log("WARN", "ALTER: " + str(e))

        conn.commit()

        # Import des produits si table vide
        cur.execute("SELECT COUNT(*) AS n FROM produits")
        row = cur.fetchone()
        if (row.get("n") or 0) == 0:
            try:
                from openpyxl import load_workbook
                wb = load_workbook("Gestion_stock_corrige.xlsx", data_only=False)
                ws = wb["Produits"]
                importes = 0
                for r in range(5, 105):
                    try:
                        sku = ws.cell(r, 1).value
                        nom = ws.cell(r, 2).value
                        if not sku or not nom:
                            continue
                        cur.execute(
                            "INSERT INTO produits (sku, nom, categorie, prix_achat, prix_vente, stock_initial, seuil) "
                            "VALUES (%s, %s, %s, %s, %s, %s, %s) ON CONFLICT (sku) DO NOTHING",
                            (str(sku), str(nom), ws.cell(r, 11).value or "",
                             int(ws.cell(r, 3).value or 0), int(ws.cell(r, 4).value or 0),
                             int(ws.cell(r, 5).value or 0), int(ws.cell(r, 10).value or 2))
                        )
                        importes += 1
                    except Exception:
                        continue
                conn.commit()
                log("DB", "{} produits importes".format(importes))
            except Exception as e:
                log("WARN", "Import Excel: " + str(e))

        cur.close()
        conn.close()
        log("DB", "Base initialisee")
    except Exception as e:
        log("ERR", "init_database: " + str(e))


try:
    init_database()
except Exception as e:
    log("ERR", "init_database global: " + str(e))


# ====================================================================
# LANCEMENT
# ====================================================================
if __name__ == "__main__":
    print("=" * 60)
    print("  Bot Messenger WeloobeAI — IA Groq robuste")
    print("  IA : {}".format("connectee" if client_ia else "NON"))
    print("  DB : {}".format("connectee" if DATABASE_URL else "NON"))
    print("=" * 60)
    app.run(host="0.0.0.0", port=5000, debug=False)