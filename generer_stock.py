# -*- coding: utf-8 -*-
"""Gestion_stock_corrige.xlsx — Version COMPLETE compatible Excel 2016.
   15 feuilles. Aucune fonction post-2016.
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

MOT_DE_PASSE = "weloobe"

# ==================== DONNÉES ====================
PRODUITS = [
    ("PC-DELL-LAT5420","Dell — Latitude 5420 i5/16Go/512Go",365000,485000,4,2,"PC portable"),
    ("PC-DELL-LAT7430","Dell — Latitude 7430 i7/16Go/512Go",520000,675000,2,2,"PC portable"),
    ("PC-HP-EL840G8","HP — EliteBook 840 G8 i5/16Go/512Go",400000,535000,3,2,"PC portable"),
    ("PC-HP-PR450G9","HP — ProBook 450 G9 i5/8Go/512Go",350000,465000,3,2,"PC portable"),
    ("PC-LEN-T14G2","Lenovo — ThinkPad T14 G2 i5/16Go/512Go",390000,520000,2,2,"PC portable"),
    ("PC-APPLE-MBA13","Apple — MacBook Air M1 8Go/256Go",510000,650000,1,2,"PC portable"),
    ("TOUR-DELL-OP7090","Dell — OptiPlex 7090 i5/16Go/512Go",340000,455000,3,2,"PC fixe"),
    ("TOUR-HP-PR400G7","HP — ProDesk 400 G7 i5/8Go/256Go",250000,345000,2,2,"PC fixe"),
    ("ECR-DELL-P2422H","Dell — P2422H 24 pouces Full HD",105000,145000,4,2,"Écran"),
    ("ECR-HP-E24G4","HP — E24 G4 24 pouces Full HD",95000,130000,2,2,"Écran"),
    ("ECR-SAMS-S24R","Samsung — S24R350 24 pouces Full HD",85000,120000,3,2,"Écran"),
    ("COMP-SSD-512","Kingston — SSD NV2 512 Go",28000,42000,8,3,"Composant"),
    ("COMP-RAM-16","Crucial — RAM DDR4 16 Go",24000,38000,6,3,"Composant"),
    ("COMP-BAT-5420","Dell — Batterie Latitude 5420",35000,55000,2,2,"Composant"),
    ("ACC-DELL-65W","Dell — Chargeur USB-C 65 W",18000,30000,5,2,"Accessoire"),
]
ACHATS = [
    (datetime.date(2026,8,2),"PC-DELL-LAT5420",5,"TechSource Cameroun"),
    (datetime.date(2026,8,5),"COMP-SSD-512",20,"Digital Pro Yaoundé"),
    (datetime.date(2026,9,3),"PC-HP-EL840G8",3,"TechSource Cameroun"),
    (datetime.date(2026,9,10),"ACC-DELL-65W",5,"Global IT Supply"),
    (datetime.date(2026,9,17),"PC-HP-EL840G8",1,"Digital Pro Yaoundé"),
]
VENTES = [
    (datetime.date(2026,8,18),"PC-DELL-LAT5420",2,"Entreprise Mboa SARL"),
    (datetime.date(2026,9,5),"ECR-DELL-P2422H",3,"Cabinet Horizon"),
    (datetime.date(2026,9,9),"PC-HP-EL840G8",2,"Entreprise Mboa SARL"),
    (datetime.date(2026,9,15),"PC-APPLE-MBA13",1,"Client comptoir"),
    (datetime.date(2026,9,20),"PC-DELL-LAT7430",1,"École Les Élites"),
    (datetime.date(2026,9,17),"PC-DELL-LAT5420",1,"Entreprise Mboa SARL"),
]
FOURNISSEURS = ["TechSource Cameroun","Digital Pro Yaoundé","Global IT Supply"]
CLIENTS = ["Entreprise Mboa SARL","Cabinet Horizon","Client comptoir","École Les Élites"]
JOURS_FR = ["Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"]

# ==================== STYLES ====================
TITRE = Font(bold=True, size=14, color="FFFFFF")
HDR   = Font(bold=True, size=11, color="FFFFFF")
BOLD  = Font(bold=True)
GRAY  = Font(color="7F7F7F", italic=True, size=10)
FILL_TITRE = PatternFill("solid", fgColor="1F4E78")
FILL_HDR   = PatternFill("solid", fgColor="2E75B6")
FILL_HDR_D = PatternFill("solid", fgColor="7F7F7F")
FILL_ALERTE = PatternFill("solid", fgColor="FCE4D6")
FILL_OK     = PatternFill("solid", fgColor="E2EFDA")
FILL_RANG   = PatternFill("solid", fgColor="F2F2F2")
FILL_WKD    = PatternFill("solid", fgColor="F8F9FA")
FILL_INPUT  = PatternFill("solid", fgColor="FFF2CC")
FILL_OR     = PatternFill("solid", fgColor="FFE699")
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT  = Alignment(horizontal="right", vertical="center")
BORDER = Border(*[Side(style="thin", color="BFBFBF")]*4)

def entete(ws, row, c1, c2, fill=FILL_HDR):
    for c in range(c1, c2+1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR; cell.fill = fill
        cell.alignment = CENTER; cell.border = BORDER

def titre(ws, cell, texte):
    ws[cell] = texte
    ws[cell].font = TITRE; ws[cell].fill = FILL_TITRE
    ws[cell].alignment = Alignment(vertical="center")

def proteger(ws, cols_libres, derniere_ligne=104):
    for r in range(1, derniere_ligne + 1):
        for c in range(1, 27):
            ws.cell(row=r, column=c).protection = \
                ws.cell(row=r, column=c).protection.copy(locked=True)
    for r in range(1, derniere_ligne + 1):
        for c in cols_libres:
            ws.cell(row=r, column=c).protection = \
                ws.cell(row=r, column=c).protection.copy(locked=False)
    ws.protection = SheetProtection(
        sheet=True, password=MOT_DE_PASSE,
        formatCells=False, formatColumns=False, formatRows=False,
        insertColumns=True, insertRows=True, deleteColumns=True, deleteRows=True,
        selectLockedCells=False, selectUnlockedCells=False,
        sort=True, autoFilter=True,
    )

wb = Workbook()

# ====================================================================
# 1) LISTES (libre)
# ====================================================================
ws = wb.active; ws.title = "Listes"
titre(ws, "A1", "LISTES DE RÉFÉRENCE")
ws["A2"] = "Feuille libre : ajoutez, modifiez ou supprimez les valeurs."
ws["A2"].font = Font(italic=True, color="595959")
ws["A3"]="SKU"; ws["B3"]="Fournisseurs"; ws["C3"]="Clients"; ws["D3"]="Produits"
entete(ws, 3, 1, 4)
for i,p in enumerate(PRODUITS, start=4):
    ws.cell(row=i, column=1, value=p[0])
    ws.cell(row=i, column=4, value=p[1])
for i,f in enumerate(FOURNISSEURS, start=4):
    ws.cell(row=i, column=2, value=f)
for i,c in enumerate(CLIENTS, start=4):
    ws.cell(row=i, column=3, value=c)
for col,w in zip("ABCD",[22,25,25,45]):
    ws.column_dimensions[col].width = w

# ====================================================================
# 2) PRODUITS
# ====================================================================
ws = wb.create_sheet("Produits")
titre(ws, "A1", "PRODUITS")
ws["A2"] = "Saisie libre : SKU, Produit, Prix, Stock, Seuil, Catégorie."
ws["A2"].font = Font(italic=True)
for c,h in enumerate(["SKU","Produit","Prix achat","Prix vente","Stock initial",
                       "Entrées","Ventes","Stock actuel","Alerte","Seuil","Catégorie",
                       "Rang alerte"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 12)
for i,p in enumerate(PRODUITS, start=5):
    ws.cell(row=i, column=1, value=p[0])
    ws.cell(row=i, column=2, value=p[1])
    ws.cell(row=i, column=3, value=p[2])
    ws.cell(row=i, column=4, value=p[3])
    ws.cell(row=i, column=5, value=p[4])
    ws.cell(row=i, column=6, value='=SUMIF(Achats!$C$5:$C$104,$A{r},Achats!$D$5:$D$104)'.format(r=i))
    ws.cell(row=i, column=7, value='=SUMIF(Ventes!$C$5:$C$104,$A{r},Ventes!$D$5:$D$104)'.format(r=i))
    ws.cell(row=i, column=8, value='=E{r}+F{r}-G{r}'.format(r=i))
    ws.cell(row=i, column=9, value='=IF(H{r}<=J{r},"A COMMANDER","OK")'.format(r=i))
    ws.cell(row=i, column=10, value=p[5])
    ws.cell(row=i, column=11, value=p[6])
    ws.cell(row=i, column=12, value='=IF(I{r}="A COMMANDER",COUNTIF($I$5:I{r},"A COMMANDER"),"")'.format(r=i))
    for c in range(1, 13):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=3).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=4).number_format = '#,##0 "FCFA"'
    for c in (9,10,12):
        ws.cell(row=i, column=c).alignment = CENTER
    ws.cell(row=i, column=12).font = GRAY
    ws.cell(row=i, column=12).fill = FILL_RANG
for col,w in zip("ABCDEFGHIJKL",[18,42,12,12,12,10,10,12,15,8,14,10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
ws.conditional_formatting.add("I5:I19",
    CellIsRule(operator="equal", formula=['"A COMMANDER"'], fill=FILL_ALERTE))
ws.conditional_formatting.add("I5:I19",
    CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK))
proteger(ws, cols_libres=[1,2,3,4,5,10,11], derniere_ligne=104)

# ====================================================================
# 3) ACHATS
# ====================================================================
ws = wb.create_sheet("Achats")
ws["A2"] = "Saisie : Date, Produit, Quantité, Fournisseur."
ws["A2"].font = Font(italic=True)
for c,h in enumerate(["Date","Produit","SKU (auto)","Quantité","Fournisseur","Rang"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 5)
ws.cell(row=4, column=6).fill = FILL_HDR_D
ws.cell(row=4, column=6).font = HDR
ws.cell(row=4, column=6).alignment = CENTER
ws.cell(row=4, column=6).border = BORDER
for i,(d,sku,q,f) in enumerate(ACHATS, start=5):
    nom = next(p[1] for p in PRODUITS if p[0]==sku)
    ws.cell(row=i, column=1, value=d).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=2, value=nom)
    ws.cell(row=i, column=3, value='=IFERROR(INDEX(Produits!$A:$A,MATCH($B{r},Produits!$B:$B,0)),"")'.format(r=i))
    ws.cell(row=i, column=4, value=q)
    ws.cell(row=i, column=5, value=f)
    ws.cell(row=i, column=6, value='=IF($A{r}="","",COUNT($A$5:A{r}))'.format(r=i))
for i in range(5, 105):
    if i > 4+len(ACHATS):
        ws.cell(row=i, column=3, value='=IFERROR(INDEX(Produits!$A:$A,MATCH($B{r},Produits!$B:$B,0)),"")'.format(r=i))
        ws.cell(row=i, column=6, value='=IF($A{r}="","",COUNT($A$5:A{r}))'.format(r=i))
    for c in range(1,7):
        ws.cell(row=i, column=c).border = BORDER
    if ws.cell(row=i, column=1).value is not None:
        ws.cell(row=i, column=1).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=6).font = GRAY
    ws.cell(row=i, column=6).fill = FILL_RANG
    ws.cell(row=i, column=6).alignment = CENTER
for col,w in zip("ABCDEF",[14,42,18,10,25,8]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
proteger(ws, cols_libres=[1,2,4,5], derniere_ligne=104)
dv1 = DataValidation(type="list", formula1="=ListeProduits", allow_blank=True)
dv2 = DataValidation(type="list", formula1="=ListeFournisseurs", allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add("B5:B104"); dv2.add("E5:E104")

# ====================================================================
# 4) VENTES
# ====================================================================
ws = wb.create_sheet("Ventes")
ws["A2"] = "Saisie : Date, Produit, Quantité, Client."
ws["A2"].font = Font(italic=True)
for c,h in enumerate(["Date","Produit","SKU (auto)","Quantité","Client","CA","Bénéfice","Rang","Clé jour"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 7)
for c in (8,9):
    ws.cell(row=4, column=c).fill = FILL_HDR_D
    ws.cell(row=4, column=c).font = HDR
    ws.cell(row=4, column=c).alignment = CENTER
    ws.cell(row=4, column=c).border = BORDER
for i,(d,sku,q,cl) in enumerate(VENTES, start=5):
    nom = next(p[1] for p in PRODUITS if p[0]==sku)
    ws.cell(row=i, column=1, value=d).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=2, value=nom)
    ws.cell(row=i, column=3, value='=IFERROR(INDEX(Produits!$A:$A,MATCH($B{r},Produits!$B:$B,0)),"")'.format(r=i))
    ws.cell(row=i, column=4, value=q)
    ws.cell(row=i, column=5, value=cl)
    ws.cell(row=i, column=6, value='=IFERROR(VLOOKUP($C{r},Produits!$A$5:$D$104,4,0)*D{r},0)'.format(r=i))
    ws.cell(row=i, column=7, value='=IFERROR((VLOOKUP($C{r},Produits!$A$5:$D$104,4,0)-VLOOKUP($C{r},Produits!$A$5:$C$104,3,0))*D{r},0)'.format(r=i))
    ws.cell(row=i, column=8, value='=IF($A{r}="","",COUNT($A$5:A{r}))'.format(r=i))
    ws.cell(row=i, column=9, value='=IF($A{r}="","",TEXT($A{r},"YYYY-MM-DD")&"-"&COUNTIF($A$5:$A{r},$A{r}))'.format(r=i))
for i in range(5, 105):
    if i > 4+len(VENTES):
        ws.cell(row=i, column=3, value='=IFERROR(INDEX(Produits!$A:$A,MATCH($B{r},Produits!$B:$B,0)),"")'.format(r=i))
        ws.cell(row=i, column=6, value='=IFERROR(VLOOKUP($C{r},Produits!$A$5:$D$104,4,0)*D{r},0)'.format(r=i))
        ws.cell(row=i, column=7, value='=IFERROR((VLOOKUP($C{r},Produits!$A$5:$D$104,4,0)-VLOOKUP($C{r},Produits!$A$5:$C$104,3,0))*D{r},0)'.format(r=i))
        ws.cell(row=i, column=8, value='=IF($A{r}="","",COUNT($A$5:A{r}))'.format(r=i))
        ws.cell(row=i, column=9, value='=IF($A{r}="","",TEXT($A{r},"YYYY-MM-DD")&"-"&COUNTIF($A$5:$A{r},$A{r}))'.format(r=i))
    for c in range(1,10):
        ws.cell(row=i, column=c).border = BORDER
    if ws.cell(row=i, column=1).value is not None:
        ws.cell(row=i, column=1).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=6).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=7).number_format = '#,##0 "FCFA"'
    for c in (8,9):
        ws.cell(row=i, column=c).font = GRAY
        ws.cell(row=i, column=c).fill = FILL_RANG
        ws.cell(row=i, column=c).alignment = CENTER
for col,w in zip("ABCDEFGHI",[14,42,18,10,25,14,14,8,16]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
proteger(ws, cols_libres=[1,2,4,5], derniere_ligne=104)
dv3 = DataValidation(type="list", formula1="=ListeProduits", allow_blank=True)
dv4 = DataValidation(type="list", formula1="=ListeClients", allow_blank=True)
ws.add_data_validation(dv3); ws.add_data_validation(dv4)
dv3.add("B5:B104"); dv4.add("E5:E104")

# ====================================================================
# 5) CALENDRIER 2026
# ====================================================================
ws = wb.create_sheet("Calendrier 2026")
titre(ws, "A1", "CALENDRIER 2026 — Synthèse journalière")
ws.merge_cells("A1:G1")
ws["A2"] = "Feuille automatique — consultable et modifiable."
ws["A2"].font = Font(italic=True, color="595959")
for c,h in enumerate(["Date","Jour","Nb ventes","CA du jour","Bénéfice du jour",
                       "Nb achats","Qté achetée"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 7)
start = datetime.date(2026, 1, 1)
for k in range(365):
    d = start + datetime.timedelta(days=k)
    i = 5 + k
    ws.cell(row=i, column=1, value=d).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=2, value=JOURS_FR[d.weekday()])
    ws.cell(row=i, column=3, value='=COUNTIF(Ventes!$A$5:$A$104,$A{r})'.format(r=i))
    ws.cell(row=i, column=4, value='=SUMPRODUCT((Ventes!$A$5:$A$104=$A{r})*Ventes!$F$5:$F$104)'.format(r=i))
    ws.cell(row=i, column=5, value='=SUMPRODUCT((Ventes!$A$5:$A$104=$A{r})*Ventes!$G$5:$G$104)'.format(r=i))
    ws.cell(row=i, column=6, value='=COUNTIF(Achats!$A$5:$A$104,$A{r})'.format(r=i))
    ws.cell(row=i, column=7, value='=SUMIF(Achats!$A$5:$A$104,$A{r},Achats!$D$5:$D$104)'.format(r=i))
    for c in range(1,8):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=4).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=5).number_format = '#,##0 "FCFA"'
    for c in (3,6,7):
        ws.cell(row=i, column=c).alignment = CENTER
    if d.weekday() >= 5:
        for c in range(1,8):
            ws.cell(row=i, column=c).fill = FILL_WKD
for col,w in zip("ABCDEFG",[14,12,12,16,16,12,12]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# ====================================================================
# 6) VUE DU JOUR
# ====================================================================
ws = wb.create_sheet("Vue du jour")
titre(ws, "A1", "VUE DU JOUR")
ws.merge_cells("A1:F1")
ws["A2"] = "Choisissez une date en B3."
ws["A2"].font = Font(italic=True, color="595959")
ws["A3"] = "Date :"; ws["A3"].font = BOLD; ws["A3"].alignment = RIGHT
ws["B3"] = datetime.date.today()
ws["B3"].number_format = 'dddd dd mmmm yyyy'
ws["B3"].font = Font(bold=True, color="1F4E78")
ws["B3"].alignment = CENTER; ws["B3"].fill = FILL_INPUT; ws["B3"].border = BORDER
for i,(lab, formule, fmt) in enumerate([
    ("Nb ventes",  '=COUNTIF(Ventes!$A$5:$A$104,$B$3)', '0'),
    ("CA du jour", '=SUMPRODUCT((Ventes!$A$5:$A$104=$B$3)*Ventes!$F$5:$F$104)', '#,##0 "FCFA"'),
    ("Bénéfice",   '=SUMPRODUCT((Ventes!$A$5:$A$104=$B$3)*Ventes!$G$5:$G$104)', '#,##0 "FCFA"'),
    ("Nb achats",  '=COUNTIF(Achats!$A$5:$A$104,$B$3)', '0'),
], start=5):
    ws.cell(row=i, column=1, value=lab).font = BOLD
    c = ws.cell(row=i, column=2, value=formule)
    c.number_format = fmt; c.border = BORDER; c.alignment = CENTER
ws["A10"] = "VENTES DU JOUR"
ws.merge_cells("A10:F10")
ws["A10"].font = HDR; ws["A10"].fill = FILL_HDR; ws["A10"].alignment = CENTER
for c,h in enumerate(["SKU","Produit","Qté","Client","CA","Bénéfice"], start=1):
    ws.cell(row=11, column=c, value=h)
entete(ws, 11, 1, 6)
for k in range(1, 21):
    i = 11 + k
    key = 'TEXT($B$3,"YYYY-MM-DD")&"-"&{k}'.format(k=k)
    ws.cell(row=i, column=1, value='=IFERROR(INDEX(Ventes!$C$5:$C$104,MATCH({key},Ventes!$I$5:$I$104,0)),"")'.format(key=key))
    ws.cell(row=i, column=2, value='=IF($A{r}="","",INDEX(Produits!$B:$B,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=3, value='=IF($A{r}="","",INDEX(Ventes!$D$5:$D$104,MATCH({key},Ventes!$I$5:$I$104,0)))'.format(r=i, key=key))
    ws.cell(row=i, column=4, value='=IF($A{r}="","",INDEX(Ventes!$E$5:$E$104,MATCH({key},Ventes!$I$5:$I$104,0)))'.format(r=i, key=key))
    ws.cell(row=i, column=5, value='=IF($A{r}="","",INDEX(Ventes!$F$5:$F$104,MATCH({key},Ventes!$I$5:$I$104,0)))'.format(r=i, key=key))
    ws.cell(row=i, column=6, value='=IF($A{r}="","",INDEX(Ventes!$G$5:$G$104,MATCH({key},Ventes!$I$5:$I$104,0)))'.format(r=i, key=key))
    for c in range(1,7):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=5).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=6).number_format = '#,##0 "FCFA"'
for col,w in zip("ABCDEF",[18,42,8,22,14,14]):
    ws.column_dimensions[col].width = w
for r in range(1, 40):
    for c in range(1, 7):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws["B3"].protection = ws["B3"].protection.copy(locked=False)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    insertColumns=True, insertRows=True, deleteColumns=True, deleteRows=True,
    selectLockedCells=False, selectUnlockedCells=False,
    sort=True, autoFilter=True)

# ====================================================================
# 7) FICHE DU JOUR
# ====================================================================
ws = wb.create_sheet("Fiche du jour")
titre(ws, "A1", "FICHE DU JOUR")
ws["A2"] = "=TODAY()"; ws["A2"].number_format = 'dddd dd mmmm yyyy'
ws["A2"].font = Font(bold=True, italic=True, color="1F4E78")
for i,(lab,formule,num) in enumerate([
    ("CA du jour",           '=SUMPRODUCT((INT(Ventes!$A$5:$A$104)=TODAY())*Ventes!$F$5:$F$104)', '#,##0 "FCFA"'),
    ("Bénéfice du jour",     '=SUMPRODUCT((INT(Ventes!$A$5:$A$104)=TODAY())*Ventes!$G$5:$G$104)', '#,##0 "FCFA"'),
    ("Nb ventes du jour",    '=COUNTIFS(Ventes!$A$5:$A$104,">="&TODAY(),Ventes!$A$5:$A$104,"<"&TODAY()+1)', '0'),
    ("Nb achats du jour",    '=COUNTIFS(Achats!$A$5:$A$104,">="&TODAY(),Achats!$A$5:$A$104,"<"&TODAY()+1)', '0'),
    ("CA du mois",           '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$F$5:$F$104)', '#,##0 "FCFA"'),
    ("Bénéfice du mois",     '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$G$5:$G$104)', '#,##0 "FCFA"'),
    ("Produits à commander", '=COUNTIF(Produits!$I$5:$I$104,"A COMMANDER")', '0'),
    ("Valeur du stock",      '=SUMPRODUCT(Produits!$C$5:$C$19,Produits!$H$5:$H$19)', '#,##0 "FCFA"'),
], start=4):
    ws.cell(row=i, column=1, value=lab).font = BOLD
    c = ws.cell(row=i, column=2, value=formule)
    c.number_format = num; c.border = BORDER
ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 20
proteger(ws, cols_libres=[], derniere_ligne=20)

# ====================================================================
# 8) PRIX DYNAMIQUES
# ====================================================================
ws = wb.create_sheet("Prix dynamiques")
titre(ws, "A1", "PRIX DYNAMIQUES")
ws["A2"] = "Marge cible en D2."
ws["A2"].font = Font(italic=True)
ws["C2"] = "Marge cible :"; ws["C2"].font = BOLD
ws["D2"] = 0.25; ws["D2"].number_format = "0%"; ws["D2"].font = BOLD
ws["D2"].fill = FILL_INPUT; ws["D2"].border = BORDER
for c,h in enumerate(["SKU","Produit","Prix achat","Marge %","Prix plancher",
                       "Prix conseillé","Prix actuel","Statut"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 8)
for i in range(5, 5+len(PRODUITS)):
    ws.cell(row=i, column=1, value='=Produits!A{r}'.format(r=i))
    ws.cell(row=i, column=2, value='=Produits!B{r}'.format(r=i))
    ws.cell(row=i, column=3, value='=Produits!C{r}'.format(r=i))
    ws.cell(row=i, column=4, value='=$D$2')
    ws.cell(row=i, column=5, value='=ROUND(C{r}*(1+D{r}),0)'.format(r=i))
    ws.cell(row=i, column=6, value='=CEILING(E{r},5000)'.format(r=i))
    ws.cell(row=i, column=7, value='=Produits!D{r}'.format(r=i))
    ws.cell(row=i, column=8, value='=IF(G{r}<E{r},"SOUS LE PLANCHER",IF(G{r}<F{r},"A AJUSTER","OK"))'.format(r=i))
    for c in range(1,9):
        ws.cell(row=i, column=c).border = BORDER
    for c in [3,5,6,7]:
        ws.cell(row=i, column=c).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=4).number_format = "0%"
    ws.cell(row=i, column=8).alignment = CENTER
for col,w in zip("ABCDEFGH",[18,42,12,10,14,14,14,18]):
    ws.column_dimensions[col].width = w
for r in range(1, 25):
    for c in range(1, 9):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws["D2"].protection = ws["D2"].protection.copy(locked=False)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    insertColumns=True, insertRows=True, deleteColumns=True, deleteRows=True,
    selectLockedCells=False, selectUnlockedCells=False,
    sort=True, autoFilter=True)

# ====================================================================
# 9) ALERTES
# ====================================================================
ws = wb.create_sheet("Alertes")
titre(ws, "A1", "ALERTES DE STOCK")
ws["A2"] = "Feuille automatique."; ws["A2"].font = Font(italic=True)
for c,h in enumerate(["SKU","Produit","Stock","Seuil","Qté à commander","Priorité"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 6)
for i in range(5, 25):
    k = i - 4
    ws.cell(row=i, column=1, value='=IFERROR(INDEX(Produits!$A$5:$A$19,MATCH({k},Produits!$L$5:$L$19,0)),"")'.format(k=k))
    ws.cell(row=i, column=2, value='=IF($A{r}="","",INDEX(Produits!$B:$B,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=3, value='=IF($A{r}="","",INDEX(Produits!$H:$H,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=4, value='=IF($A{r}="","",INDEX(Produits!$J:$J,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=5, value='=IF($A{r}="","",MAX(3,D{r}*2-C{r}))'.format(r=i))
    ws.cell(row=i, column=6, value='=IF($A{r}="","",IF(C{r}=0,"URGENT",IF(C{r}<=D{r}/2,"HAUT","NORMAL")))'.format(r=i))
    for c in range(1,7):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=6).alignment = CENTER
for col,w in zip("ABCDEF",[18,42,10,8,16,12]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
proteger(ws, cols_libres=[], derniere_ligne=24)

# ====================================================================
# 10) TOP PRODUITS (tri automatique compatible 2016)
# ====================================================================
ws = wb.create_sheet("Top produits")
titre(ws, "A1", "TOP PRODUITS — Classement par CA")
ws["A2"] = "Tri automatique par chiffre d'affaires décroissant."
ws["A2"].font = Font(italic=True, color="595959")
for c,h in enumerate(["Rang","SKU","Produit","Qté vendue","CA généré","Marge totale","Marge %"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 7)

# Colonnes auxiliaires I, J, K, L (calculs)
for i in range(5, 5+15):
    r = i
    ws.cell(row=r, column=9, value='=Produits!A{p}'.format(p=r))
    ws.cell(row=r, column=10, value='=SUMIF(Ventes!$C$5:$C$104,$I{r},Ventes!$F$5:$F$104)'.format(r=r))
    ws.cell(row=r, column=11, value='=SUMIF(Ventes!$C$5:$C$104,$I{r},Ventes!$D$5:$D$104)'.format(r=r))
    ws.cell(row=r, column=12, value='=SUMIF(Ventes!$C$5:$C$104,$I{r},Ventes!$G$5:$G$104)'.format(r=r))

# Colonnes résultats A-G
for i in range(5, 5+15):
    r = i
    k = r - 4
    ws.cell(row=r, column=1, value=k)
    ws.cell(row=r, column=5, value='=LARGE($J$5:$J$19,{k})'.format(k=k))
    ws.cell(row=r, column=2, value='=IFERROR(INDEX($I$5:$I$19,MATCH($E{r},$J$5:$J$19,0)),"")'.format(r=r))
    ws.cell(row=r, column=3, value='=IFERROR(INDEX(Produits!$B:$B,MATCH($B{r},Produits!$A:$A,0)),"")'.format(r=r))
    ws.cell(row=r, column=4, value='=IFERROR(INDEX($K$5:$K$19,MATCH($E{r},$J$5:$J$19,0)),0)'.format(r=r))
    ws.cell(row=r, column=6, value='=IFERROR(INDEX($L$5:$L$19,MATCH($E{r},$J$5:$J$19,0)),0)'.format(r=r))
    ws.cell(row=r, column=7, value='=IF(OR($E{r}=0,$E{r}=""),"",$F{r}/$E{r})'.format(r=r))
    for c in range(1,8):
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=5).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=6).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=7).number_format = '0.0%'
    for c in (1,4):
        ws.cell(row=r, column=c).alignment = CENTER

ws.column_dimensions['I'].hidden = True
ws.column_dimensions['J'].hidden = True
ws.column_dimensions['K'].hidden = True
ws.column_dimensions['L'].hidden = True
for col,w in zip("ABCDEFG",[8,18,42,12,14,14,10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
for r in range(1, 25):
    for c in range(1, 15):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    selectLockedCells=False, selectUnlockedCells=False, sort=True, autoFilter=True)

# ====================================================================
# 11) CLIENTS (compatible 2016 — SUMPRODUCT(MAX) au lieu de MAXIFS)
# ====================================================================
ws = wb.create_sheet("Clients")
titre(ws, "A1", "CLIENTS — Historique")
ws["A2"] = "CA, nombre d'achats et dernière commande par client."
ws["A2"].font = Font(italic=True, color="595959")
for c,h in enumerate(["Client","Nb commandes","CA total","Panier moyen","Dernière commande"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 5)
for idx, nom in enumerate(CLIENTS):
    r = 5 + idx
    ws.cell(row=r, column=1, value=nom)
    ws.cell(row=r, column=2, value='=COUNTIF(Ventes!$E$5:$E$104,$A{r})'.format(r=r))
    ws.cell(row=r, column=3, value='=SUMIF(Ventes!$E$5:$E$104,$A{r},Ventes!$F$5:$F$104)'.format(r=r))
    ws.cell(row=r, column=4, value='=IF($B{r}=0,0,$C{r}/$B{r})'.format(r=r))
    # Compatible 2016 : SUMPRODUCT(MAX(...)) au lieu de MAXIFS
    ws.cell(row=r, column=5,
        value='=IF($B{r}=0,"",SUMPRODUCT(MAX((Ventes!$E$5:$E$104=$A{r})*(Ventes!$A$5:$A$104))))'.format(r=r))
    for c in range(1,6):
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=3).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=4).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=5).number_format = 'yyyy-mm-dd'
    for c in (2,5):
        ws.cell(row=r, column=c).alignment = CENTER
for col,w in zip("ABCDE",[30,14,16,16,18]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
for r in range(1, 25):
    for c in range(1, 6):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    selectLockedCells=False, selectUnlockedCells=False, sort=True, autoFilter=True)

# ====================================================================
# 12) RAPPORT MENSUEL
# ====================================================================
ws = wb.create_sheet("Rapport mensuel")
titre(ws, "A1", "RAPPORT MENSUEL 2026")
ws["A2"] = "Indicateurs consolidés mois par mois."
ws["A2"].font = Font(italic=True, color="595959")
for c,h in enumerate(["Mois","Nb ventes","CA","Bénéfice","Marge %","Nb achats","Qté achetée"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 7)
mois_fr = ["Janvier","Février","Mars","Avril","Mai","Juin",
           "Juillet","Août","Septembre","Octobre","Novembre","Décembre"]
for idx, mois in enumerate(mois_fr):
    r = 5 + idx
    m = idx + 1
    ws.cell(row=r, column=1, value=mois).font = BOLD
    ws.cell(row=r, column=2, value='=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")="2026-{m:02d}")*1)'.format(m=m))
    ws.cell(row=r, column=3, value='=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")="2026-{m:02d}")*Ventes!$F$5:$F$104)'.format(m=m))
    ws.cell(row=r, column=4, value='=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")="2026-{m:02d}")*Ventes!$G$5:$G$104)'.format(m=m))
    ws.cell(row=r, column=5, value='=IF($C{r}=0,"",$D{r}/$C{r})'.format(r=r))
    ws.cell(row=r, column=6, value='=SUMPRODUCT((TEXT(Achats!$A$5:$A$104,"YYYY-MM")="2026-{m:02d}")*1)'.format(m=m))
    ws.cell(row=r, column=7, value='=SUMPRODUCT((TEXT(Achats!$A$5:$A$104,"YYYY-MM")="2026-{m:02d}")*Achats!$D$5:$D$104)'.format(m=m))
    for c in range(1,8):
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=3).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=4).number_format = '#,##0 "FCFA"'
    ws.cell(row=r, column=5).number_format = '0.0%'
    for c in (2,6,7):
        ws.cell(row=r, column=c).alignment = CENTER
# Ligne TOTAL
r_tot = 17
ws.cell(row=r_tot, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=r_tot, column=2, value='=SUM(B5:B16)')
ws.cell(row=r_tot, column=3, value='=SUM(C5:C16)')
ws.cell(row=r_tot, column=4, value='=SUM(D5:D16)')
ws.cell(row=r_tot, column=5, value='=IF(C{t}=0,"",D{t}/C{t})'.format(t=r_tot))
ws.cell(row=r_tot, column=6, value='=SUM(F5:F16)')
ws.cell(row=r_tot, column=7, value='=SUM(G5:G16)')
for c in range(1,8):
    ws.cell(row=r_tot, column=c).border = BORDER
    ws.cell(row=r_tot, column=c).fill = FILL_OR
ws.cell(row=r_tot, column=3).number_format = '#,##0 "FCFA"'
ws.cell(row=r_tot, column=4).number_format = '#,##0 "FCFA"'
ws.cell(row=r_tot, column=5).number_format = '0.0%'
for col,w in zip("ABCDEFG",[14,12,16,16,10,12,12]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
for r in range(1, 25):
    for c in range(1, 8):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    selectLockedCells=False, selectUnlockedCells=False, sort=True, autoFilter=True)

# ====================================================================
# 13) PRÉVISIONS (compatible 2016)
# ====================================================================
ws = wb.create_sheet("Prévisions")
titre(ws, "A1", "PRÉVISIONS DE RUPTURE")
ws["A2"] = "Délai estimé avant rupture selon la vitesse de vente."
ws["A2"].font = Font(italic=True, color="595959")
for c,h in enumerate(["SKU","Produit","Stock actuel","Qté vendue (30j)","Vitesse/jour","Jours avant rupture","Statut"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 7)
for i in range(5, 5+15):
    r = i
    ws.cell(row=r, column=1, value='=Produits!A{p}'.format(p=r))
    ws.cell(row=r, column=2, value='=Produits!B{p}'.format(p=r))
    ws.cell(row=r, column=3, value='=Produits!H{p}'.format(p=r))
    # Qté vendue sur 30 jours (compatible 2016)
    ws.cell(row=r, column=4,
        value='=SUMPRODUCT((Ventes!$C$5:$C$104=$A{r})*(Ventes!$A$5:$A$104>=TODAY()-30)*Ventes!$D$5:$D$104)'.format(r=r))
    ws.cell(row=r, column=5, value='=IF($D{r}=0,0,$D{r}/30)'.format(r=r))
    ws.cell(row=r, column=6, value='=IF($E{r}=0,"",IF($C{r}<=0,0,ROUND($C{r}/$E{r},0)))'.format(r=r))
    ws.cell(row=r, column=7,
        value='=IF($C{r}<=0,"RUPTURE",IF(AND(ISNUMBER($F{r}),$F{r}<=7),"CRITIQUE",IF(AND(ISNUMBER($F{r}),$F{r}<=30),"A SURVEILLER","OK")))'.format(r=r))
    for c in range(1,8):
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=5).number_format = '0.0'
    for c in (3,4,6,7):
        ws.cell(row=r, column=c).alignment = CENTER
for col,w in zip("ABCDEFG",[18,42,12,16,12,16,14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
# Mise en forme conditionnelle (texte exact)
ws.conditional_formatting.add("G5:G19",
    CellIsRule(operator="equal", formula=['"RUPTURE"'], fill=FILL_ALERTE))
ws.conditional_formatting.add("G5:G19",
    CellIsRule(operator="equal", formula=['"CRITIQUE"'], fill=FILL_ALERTE))
ws.conditional_formatting.add("G5:G19",
    CellIsRule(operator="equal", formula=['"A SURVEILLER"'], fill=FILL_OR))
ws.conditional_formatting.add("G5:G19",
    CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK))
for r in range(1, 25):
    for c in range(1, 8):
        ws.cell(row=r, column=c).protection = ws.cell(row=r, column=c).protection.copy(locked=True)
ws.protection = SheetProtection(sheet=True, password=MOT_DE_PASSE,
    formatCells=False, formatColumns=False, formatRows=False,
    selectLockedCells=False, selectUnlockedCells=False, sort=True, autoFilter=True)

# ====================================================================
# 14) DONNÉES CHATBOT
# ====================================================================
ws = wb.create_sheet("Données chatbot")
titre(ws, "A1", "DONNÉES POUR LE CHATBOT")
ws["A2"] = "Feuille automatique."; ws["A2"].font = Font(italic=True)
for c,h in enumerate(["SKU","Nom","Catégorie","Prix","Stock","Dispo","Alerte","Mots-clés"], start=1):
    ws.cell(row=4, column=c, value=h)
entete(ws, 4, 1, 8)
for i in range(5, 5+len(PRODUITS)):
    ws.cell(row=i, column=1, value='=Produits!A{r}'.format(r=i))
    ws.cell(row=i, column=2, value='=Produits!B{r}'.format(r=i))
    ws.cell(row=i, column=3, value='=Produits!K{r}'.format(r=i))
    ws.cell(row=i, column=4, value='=Produits!D{r}'.format(r=i))
    ws.cell(row=i, column=5, value='=Produits!H{r}'.format(r=i))
    ws.cell(row=i, column=6, value='=IF(E{r}>0,"En stock","Rupture")'.format(r=i))
    ws.cell(row=i, column=7, value='=Produits!I{r}'.format(r=i))
    ws.cell(row=i, column=8, value='=LOWER(B{r}&" "&C{r})'.format(r=i))
    for c in range(1,9):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=4).number_format = '#,##0 "FCFA"'
for col,w in zip("ABCDEFGH",[18,42,14,14,8,10,15,55]):
    ws.column_dimensions[col].width = w
ws["J4"] = "KPI"; ws["K4"] = "Valeur"
entete(ws, 4, 10, 11)
for i,(k,formule) in enumerate([
    ("chiffre_affaires_mois", '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$F$5:$F$104)'),
    ("benefice_mois",         '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$G$5:$G$104)'),
    ("produits_en_alerte",    '=COUNTIF(Produits!$I$5:$I$104,"A COMMANDER")'),
    ("valeur_stock",          '=SUMPRODUCT(Produits!$C$5:$C$19,Produits!$H$5:$H$19)'),
    ("nb_ventes_total",       '=COUNTA(Ventes!$A$5:$A$104)'),
    ("nb_achats_total",       '=COUNTA(Achats!$A$5:$A$104)'),
], start=5):
    ws.cell(row=i, column=10, value=k).font = BOLD
    ws.cell(row=i, column=11, value=formule).number_format = '#,##0'
ws.column_dimensions["J"].width = 25
ws.column_dimensions["K"].width = 18
proteger(ws, cols_libres=[], derniere_ligne=30)

# ====================================================================
# 15) DASHBOARD (position 1) + GRAPHIQUES
# ====================================================================
ws = wb.create_sheet("Dashboard", 0)
titre(ws, "A1", "DASHBOARD — ACTIVITÉ COMMERCIALE")
ws.merge_cells("A1:O1")
ws["A2"] = "Feuille 100 % automatique."
ws["A2"].font = Font(italic=True, color="595959")
for cell, lab in [("A4","STOCK TOTAL (unités)"),("C4","CHIFFRE D'AFFAIRES"),
                  ("E4","BÉNÉFICE BRUT"),("H4","À COMMANDER")]:
    ws[cell] = lab; ws[cell].font = HDR; ws[cell].fill = FILL_HDR; ws[cell].alignment = CENTER
ws["A5"] = "=SUM(Produits!$H$5:$H$104)"
ws["C5"] = "=SUM(Ventes!$F$5:$F$104)"
ws["E5"] = "=SUM(Ventes!$G$5:$G$104)"
ws["H5"] = '=COUNTIF(Produits!$I$5:$I$104,"A COMMANDER")'
for c in ["A5","C5","E5","H5"]:
    ws[c].font = Font(bold=True, size=13, color="1F4E78")
    ws[c].alignment = CENTER; ws[c].border = BORDER
ws["C5"].number_format = '#,##0 "FCFA"'
ws["E5"].number_format = '#,##0 "FCFA"'
ws["A6"] = "dont ce mois :"; ws["A6"].font = GRAY
ws["C6"] = '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$F$5:$F$104)'
ws["E6"] = '=SUMPRODUCT((TEXT(Ventes!$A$5:$A$104,"YYYY-MM")=TEXT(TODAY(),"YYYY-MM"))*Ventes!$G$5:$G$104)'
for c in ["C6","E6"]:
    ws[c].font = GRAY; ws[c].alignment = CENTER
    ws[c].number_format = '#,##0 "FCFA"'
ws["A7"] = "PRODUITS A COMMANDER"; ws["F7"] = "10 DERNIERES VENTES"; ws["L7"] = "10 DERNIERS ACHATS"
for c in ["A7","F7","L7"]:
    ws[c].font = HDR; ws[c].fill = FILL_HDR; ws[c].alignment = CENTER
ws.merge_cells("A7:D7"); ws.merge_cells("F7:J7"); ws.merge_cells("L7:N7")
for i,h in enumerate(["SKU","Produit","Stock","Seuil"], start=1):
    ws.cell(row=8, column=i, value=h)
for i,h in enumerate(["Date","SKU","Client","Qté","CA"], start=6):
    ws.cell(row=8, column=i, value=h)
for i,h in enumerate(["Date","SKU","Qté"], start=12):
    ws.cell(row=8, column=i, value=h)
entete(ws, 8, 1, 4); entete(ws, 8, 6, 10); entete(ws, 8, 12, 14)
for i in range(9, 19):
    k = i - 8
    ws.cell(row=i, column=1, value='=IFERROR(INDEX(Produits!$A$5:$A$19,MATCH({k},Produits!$L$5:$L$19,0)),"")'.format(k=k))
    ws.cell(row=i, column=2, value='=IF($A{r}="","",INDEX(Produits!$B:$B,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=3, value='=IF($A{r}="","",INDEX(Produits!$H:$H,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    ws.cell(row=i, column=4, value='=IF($A{r}="","",INDEX(Produits!$J:$J,MATCH($A{r},Produits!$A:$A,0)))'.format(r=i))
    rk_v = 'COUNT(Ventes!$H$5:$H$104)-ROWS($F$9:F{r})+1'.format(r=i)
    ws.cell(row=i, column=6, value='=IFERROR(INDEX(Ventes!$A$5:$A$104,MATCH({rk},Ventes!$H$5:$H$104,0)),"")'.format(rk=rk_v))
    ws.cell(row=i, column=7, value='=IF($F{r}="","",INDEX(Ventes!$C$5:$C$104,MATCH({rk},Ventes!$H$5:$H$104,0)))'.format(r=i, rk=rk_v))
    ws.cell(row=i, column=8, value='=IF($F{r}="","",INDEX(Ventes!$E$5:$E$104,MATCH({rk},Ventes!$H$5:$H$104,0)))'.format(r=i, rk=rk_v))
    ws.cell(row=i, column=9, value='=IF($F{r}="","",INDEX(Ventes!$D$5:$D$104,MATCH({rk},Ventes!$H$5:$H$104,0)))'.format(r=i, rk=rk_v))
    ws.cell(row=i, column=10, value='=IF($F{r}="","",INDEX(Ventes!$F$5:$F$104,MATCH({rk},Ventes!$H$5:$H$104,0)))'.format(r=i, rk=rk_v))
    rk_a = 'COUNT(Achats!$F$5:$F$104)-ROWS($L$9:L{r})+1'.format(r=i)
    ws.cell(row=i, column=12, value='=IFERROR(INDEX(Achats!$A$5:$A$104,MATCH({rk},Achats!$F$5:$F$104,0)),"")'.format(rk=rk_a))
    ws.cell(row=i, column=13, value='=IF($L{r}="","",INDEX(Achats!$C$5:$C$104,MATCH({rk},Achats!$F$5:$F$104,0)))'.format(r=i, rk=rk_a))
    ws.cell(row=i, column=14, value='=IF($L{r}="","",INDEX(Achats!$D$5:$D$104,MATCH({rk},Achats!$F$5:$F$104,0)))'.format(r=i, rk=rk_a))
    for c in range(1,15):
        ws.cell(row=i, column=c).border = BORDER
    ws.cell(row=i, column=6).number_format = 'yyyy-mm-dd'
    ws.cell(row=i, column=10).number_format = '#,##0 "FCFA"'
    ws.cell(row=i, column=12).number_format = 'yyyy-mm-dd'
for col,w in zip("ABCDEFGHIJKLMNO",[18,42,10,8,2,12,18,22,8,14,2,12,18,8,2]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A9"

# Graphique 1 : évolution CA mensuel
ws_rm = wb["Rapport mensuel"]
chart1 = LineChart()
chart1.title = "Évolution du CA 2026"
chart1.height = 8; chart1.width = 16
data1 = Reference(ws_rm, min_col=3, min_row=4, max_row=16)
cats1 = Reference(ws_rm, min_col=1, min_row=5, max_row=16)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
ws.add_chart(chart1, "A21")

# Graphique 2 : Top 5 produits
ws_top = wb["Top produits"]
chart2 = BarChart()
chart2.title = "Top 5 produits par CA"
chart2.height = 8; chart2.width = 16
data2 = Reference(ws_top, min_col=5, min_row=4, max_row=9)
cats2 = Reference(ws_top, min_col=3, min_row=5, max_row=9)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws.add_chart(chart2, "J21")

# Graphique 3 : stock par catégorie (données auxiliaires)
ws["Z1"] = "Catégorie"; ws["AA1"] = "Stock"
categories = ["PC portable","PC fixe","Écran","Composant","Accessoire"]
for idx, cat in enumerate(categories):
    r = 2 + idx
    ws.cell(row=r, column=26, value=cat)
    ws.cell(row=r, column=27,
        value='=SUMIF(Produits!$K$5:$K$19,Z{r},Produits!$H$5:$H$19)'.format(r=r))

chart3 = PieChart()
chart3.title = "Stock par catégorie"
chart3.height = 8; chart3.width = 12
data3 = Reference(ws, min_col=27, min_row=1, max_row=6)
cats3 = Reference(ws, min_col=26, min_row=2, max_row=6)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws.add_chart(chart3, "A40")

# Masquer les colonnes auxiliaires
ws.column_dimensions['Z'].hidden = True
ws.column_dimensions['AA'].hidden = True

proteger(ws, cols_libres=[], derniere_ligne=60)

# ====================================================================
# PLAGES NOMMÉES FIXES (pas d'OFFSET)
# ====================================================================
for nom, ref in {
    "ListeProduits":     "Produits!$B$5:$B$104",
    "ListeFournisseurs": "Listes!$B$4:$B$100",
    "ListeClients":      "Listes!$C$4:$C$100",
    "ListeSKU":          "Produits!$A$5:$A$104",
}.items():
    try:
        wb.defined_names.add(DefinedName(nom, attr_text=ref))
    except AttributeError:
        wb.defined_names.append(DefinedName(nom, attr_text=ref))

out = "Gestion_stock_corrige.xlsx"
wb.save(out)
print("=" * 60)
print("Fichier genere : " + out)
print("15 feuilles — 100% compatible Excel 2016")
print("=" * 60)
print("Feuilles :")
for s in wb.sheetnames:
    print("  - " + s)
print("=" * 60)
print("Aucune fonction post-2016 :")
print("  - MAXIFS remplace par SUMPRODUCT(MAX(...))")
print("  - Pas d'OFFSET / XLOOKUP / IFS / TEXTJOIN")
print("  - Autocompletion native preservee")
print("Mot de passe : " + MOT_DE_PASSE)
print("=" * 60)