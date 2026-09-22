# -*- coding: utf-8 -*-
"""Initialise la base SQLite avec les produits du fichier Excel."""
import sqlite3
import datetime
from pathlib import Path
from openpyxl import load_workbook

DB = "lab.db"
XLSX = "Gestion_stock_corrige.xlsx"


def init():
    if Path(DB).exists():
        Path(DB).unlink()
        print("Ancienne base supprimée.")

    conn = sqlite3.connect(DB)
    c = conn.cursor()

    # Table produits
    c.execute("""
        CREATE TABLE produits (
            sku TEXT PRIMARY KEY,
            nom TEXT NOT NULL,
            categorie TEXT,
            prix_achat INTEGER,
            prix_vente INTEGER,
            stock_initial INTEGER,
            seuil INTEGER
        )
    """)

    # Table ventes
    c.execute("""
        CREATE TABLE ventes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT,
            sku TEXT,
            quantite INTEGER,
            client TEXT,
            canal TEXT DEFAULT 'lab'
        )
    """)

    # Table clients (Messenger)
    c.execute("""
        CREATE TABLE clients (
            psid TEXT PRIMARY KEY,
            prenom TEXT,
            nom TEXT,
            derniere_interaction TEXT,
            derniere_intention TEXT,
            contexte TEXT
        )
    """)

    # Table messages (historique conversation)
    c.execute("""
        CREATE TABLE messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            psid TEXT,
            direction TEXT,
            contenu TEXT,
            timestamp TEXT
        )
    """)

    # Table propositions (offres faites au client)
    c.execute("""
        CREATE TABLE propositions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            psid TEXT,
            sku TEXT,
            timestamp TEXT,
            statut TEXT
        )
    """)

    # Charger les produits depuis Excel
    if not Path(XLSX).exists():
        print("ATTENTION : fichier Excel introuvable. Table produits vide.")
    else:
        wb = load_workbook(XLSX, data_only=False)
        ws = wb["Produits"]
        produits = []
        for r in range(5, 105):
            sku = ws.cell(r, 1).value
            nom = ws.cell(r, 2).value
            if not sku or not nom:
                continue
            produits.append((
                str(sku),
                str(nom),
                ws.cell(r, 11).value or "",
                int(ws.cell(r, 3).value or 0),
                int(ws.cell(r, 4).value or 0),
                int(ws.cell(r, 5).value or 0),
                int(ws.cell(r, 10).value or 2)
            ))
        c.executemany("INSERT INTO produits VALUES (?, ?, ?, ?, ?, ?, ?)", produits)
        print("{} produits importés.".format(len(produits)))

    conn.commit()
    conn.close()
    print("Base {} créée.".format(DB))


if __name__ == "__main__":
    init()